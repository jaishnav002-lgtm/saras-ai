"""Extraction-only timetable service.

Every extractor returns the same ExamEntry shape.  It deliberately has no
database or calendar dependency, so a future vision fallback can implement
``extract`` here without changing review or calendar code.
"""
from __future__ import annotations

import csv
import datetime as dt
import logging
import os
import re
from dataclasses import dataclass
from pathlib import Path
from statistics import median
from typing import Iterable, Optional

from app.schemas import ExamConfidence, ExamEntry

logger = logging.getLogger(__name__)

DATE_RE = re.compile(r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{1,2}\s+(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)(?:\s+\d{2,4})?|(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\s+\d{1,2},?\s+\d{2,4})\b", re.I)
TIME_RE = re.compile(r"\b(\d{1,2}(?::\d{2})?\s*(?:AM|PM)?|\d{1,2}:\d{2})\s*(?:-|–|to)\s*(\d{1,2}(?::\d{2})?\s*(?:AM|PM)?|\d{1,2}:\d{2})\b", re.I)
VENUE_RE = re.compile(r"\b(?:venue\s*[:=-]?\s*)?(?:hall\s*[A-Za-z0-9-]+|room(?:\s*(?:no\.?|number))?\s*[:#-]?\s*[A-Za-z0-9-]+|[A-Z]-\d{2,4}|exam\s+hall)\b", re.I)


@dataclass
class OCRItem:
    text: str
    confidence: float
    x1: float
    y1: float
    x2: float
    y2: float
    page: int = 1

    @property
    def cx(self) -> float:
        return (self.x1 + self.x2) / 2

    @property
    def cy(self) -> float:
        return (self.y1 + self.y2) / 2

    @property
    def height(self) -> float:
        return max(1, self.y2 - self.y1)


_ocr = None


def _get_ocr():
    """Lazy singleton: model download/start-up happens only on image/PDF use."""
    global _ocr
    if _ocr is None:
        # PaddleX otherwise writes under the login home, which is outside the
        # project runtime on sandboxed/local installations.
        os.environ.setdefault("PADDLE_PDX_CACHE_HOME", str(Path(__file__).resolve().parents[1] / ".paddlex"))
        from paddleocr import PaddleOCR
        logger.info("Initializing CPU PaddleOCR timetable service")
        _ocr = PaddleOCR(lang="en")
    return _ocr


def _ocr_image(path: Path, page: int = 1) -> list[OCRItem]:
    engine = _get_ocr()
    result = engine.predict(str(path))
    items: list[OCRItem] = []
    for result_page in result:
        data = result_page if isinstance(result_page, dict) else getattr(result_page, "json", result_page)
        if callable(data):
            data = data()
        if isinstance(data, dict) and "res" in data:
            data = data["res"]
        if not isinstance(data, dict):
            continue
        texts = data.get("rec_texts", [])
        scores = data.get("rec_scores", [])
        boxes = data.get("rec_boxes", data.get("dt_polys", []))
        for text, score, box in zip(texts, scores, boxes):
            if not str(text).strip():
                continue
            # Paddle v3 rec_boxes are [x1,y1,x2,y2]; dt_polys are four points.
            if len(box) == 4 and not isinstance(box[0], (list, tuple)):
                x1, y1, x2, y2 = map(float, box)
            else:
                xs = [point[0] for point in box]
                ys = [point[1] for point in box]
                x1, x2, y1, y2 = min(xs), max(xs), min(ys), max(ys)
            items.append(OCRItem(str(text).strip(), float(score), x1, y1, x2, y2, page))
    return items


def _normalise_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


HEADER_KINDS = {
    "subject": ("subject", "course", "coursename", "paper", "subjectname", "coursecode"),
    "date": ("date", "examdate", "examinationdate", "daydate"),
    "time": ("time", "examtime", "timing", "starttime", "duration"),
    "venue": ("venue", "room", "roomno", "hall", "examhall", "location"),
}


def _header_map(headers: Iterable[object]) -> dict[str, int]:
    answer: dict[str, int] = {}
    for index, header in enumerate(headers):
        normal = _normalise_header(header)
        for kind, names in HEADER_KINDS.items():
            if kind not in answer and any(name == normal or name in normal for name in names):
                answer[kind] = index
    return answer


def _normalise_date(value: object) -> tuple[Optional[str], Optional[str]]:
    if isinstance(value, dt.datetime):
        return value.date().isoformat(), None
    if isinstance(value, dt.date):
        return value.isoformat(), None
    raw = str(value or "").strip()
    if not raw:
        return None, "Missing exam date."
    # Explicit dates only.  No fallback to current year.
    formats = ("%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y", "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y")
    for fmt in formats:
        try:
            return dt.datetime.strptime(raw, fmt).date().isoformat(), None
        except ValueError:
            pass
    if re.search(r"\b\d{1,2}\s+[A-Za-z]+\b", raw):
        return None, "Date has no year; confirm it in review."
    return None, f"Invalid date '{raw}'."


def _normalise_time(value: object) -> tuple[Optional[str], Optional[str], Optional[str]]:
    raw = str(value or "").strip()
    if not raw:
        return None, None, None
    match = TIME_RE.search(raw)
    if not match:
        return None, None, "Could not read exam time."
    def parse(part: str) -> Optional[dt.time]:
        part = re.sub(r"\s+", " ", part.strip().upper())
        for fmt in ("%I:%M %p", "%I %p", "%H:%M"):
            try: return dt.datetime.strptime(part, fmt).time()
            except ValueError: pass
        return None
    start, end = parse(match.group(1)), parse(match.group(2))
    if not start or not end:
        return None, None, "Could not normalize exam time."
    if end <= start:
        return start.strftime("%H:%M"), end.strftime("%H:%M"), "End time must be after start time."
    return start.strftime("%H:%M"), end.strftime("%H:%M"), None


def _entry(subject: object, date: object, timing: object = "", venue: object = "", confidence: float = 1.0, page: Optional[int] = None) -> ExamEntry:
    warnings: list[str] = []
    subject_text = str(subject or "").strip()
    if not subject_text or subject_text == "-":
        warnings.append("Missing subject.")
    iso_date, date_warning = _normalise_date(date)
    if date_warning: warnings.append(date_warning)
    start, end, time_warning = _normalise_time(timing)
    if time_warning: warnings.append(time_warning)
    venue_match = VENUE_RE.search(str(venue or ""))
    return ExamEntry(subject=subject_text, date=iso_date, start_time=start, end_time=end,
                     venue=venue_match.group(0) if venue_match else (str(venue).strip() or None),
                     confidence=ExamConfidence(subject=confidence, date=confidence if iso_date else 0, overall=confidence),
                     warnings=warnings, source_page=page)


def parse_spreadsheet_rows(rows: list[list[object]]) -> list[ExamEntry]:
    if not rows:
        return []
    headers = _header_map(rows[0])
    if "subject" not in headers or "date" not in headers:
        return [ExamEntry(warnings=["Could not identify Subject and Date columns. Add exams manually or correct the sheet headers."])]
    entries = []
    for row in rows[1:]:
        if not any(str(cell or "").strip() for cell in row):
            continue
        get = lambda kind: row[headers[kind]] if kind in headers and headers[kind] < len(row) else ""
        entries.append(_entry(get("subject"), get("date"), get("time"), get("venue")))
    return entries


def _cluster_rows(items: list[OCRItem]) -> list[list[OCRItem]]:
    if not items: return []
    threshold = max(8, median(item.height for item in items) * 0.8)
    rows: list[list[OCRItem]] = []
    for item in sorted(items, key=lambda item: item.cy):
        if not rows or abs(item.cy - median(x.cy for x in rows[-1])) > threshold:
            rows.append([item])
        else:
            rows[-1].append(item)
    return [sorted(row, key=lambda item: item.cx) for row in rows]


def parse_ocr_grid(items: list[OCRItem]) -> list[ExamEntry]:
    """Use y/x coordinates to reconstruct rows and header-defined cells."""
    rows = _cluster_rows(items)
    header_index = next((i for i, row in enumerate(rows) if {k for k in _header_map([x.text for x in row])} >= {"subject", "date"}), None)
    if header_index is None:
        return _parse_ocr_text_blocks(items)
    header = rows[header_index]
    mapping = _header_map([x.text for x in header])
    columns = {kind: header[index].cx for kind, index in mapping.items()}
    boundaries = sorted((x, kind) for kind, x in columns.items())
    entries: list[ExamEntry] = []
    for row in rows[header_index + 1:]:
        cells: dict[str, list[OCRItem]] = {kind: [] for kind in columns}
        for item in row:
            kind = min(columns, key=lambda candidate: abs(item.cx - columns[candidate]))
            cells[kind].append(item)
        text = {kind: " ".join(x.text for x in sorted(cell, key=lambda x: x.cx)) for kind, cell in cells.items()}
        if text.get("subject") or text.get("date"):
            confidence = median([x.confidence for cell in cells.values() for x in cell] or [0.0])
            entries.append(_entry(text.get("subject", ""), text.get("date", ""), text.get("time", ""), text.get("venue", ""), confidence, row[0].page))
    return entries


def _parse_ocr_text_blocks(items: list[OCRItem]) -> list[ExamEntry]:
    # Spatially reconstructed lines remain useful for non-standard tables.
    lines = [(" ".join(item.text for item in row), row) for row in _cluster_rows(items)]
    entries: list[ExamEntry] = []
    for index, (line, row) in enumerate(lines):
        date_match = DATE_RE.search(line)
        if not date_match:
            continue
        context = " ".join(text for text, _ in lines[max(0, index - 1): index + 2])
        subject = re.sub(DATE_RE, "", line).strip(" -|:")
        subject = TIME_RE.sub("", subject).strip(" -|:") or (lines[index - 1][0] if index else "")
        time_match = TIME_RE.search(context)
        venue_match = VENUE_RE.search(context)
        confidence = median([item.confidence for item in row])
        entries.append(_entry(subject, date_match.group(0), time_match.group(0) if time_match else "", venue_match.group(0) if venue_match else "", confidence, row[0].page))
    return entries


def _deduplicate(entries: list[ExamEntry]) -> tuple[list[ExamEntry], int]:
    """Repeated PDF rows are not silently imported twice."""
    seen: set[tuple[str, Optional[str], Optional[str]]] = set()
    unique: list[ExamEntry] = []
    duplicates = 0
    for entry in entries:
        key = (_normalise_header(entry.subject), entry.date, entry.start_time)
        if entry.date and entry.subject and key in seen:
            duplicates += 1
            continue
        seen.add(key)
        unique.append(entry)
    return unique, duplicates


def extract_exams_from_file(path: Path, extension: str) -> tuple[list[ExamEntry], list[str], int]:
    extension = extension.lower()
    if extension == "csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return parse_spreadsheet_rows(list(csv.reader(handle))), [], 0
    if extension == "xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = [list(row) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)]
        return parse_spreadsheet_rows(rows), [], 0
    items: list[OCRItem] = []
    if extension == "pdf":
        import fitz
        document = fitz.open(path)
        try:
            for page_no, page in enumerate(document, 1):
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image_path = path.with_name(f"{path.stem}-page-{page_no}.png")
                pix.save(str(image_path))
                items.extend(_ocr_image(image_path, page_no))
                image_path.unlink(missing_ok=True)
        finally:
            document.close()
    else:
        items = _ocr_image(path)
    entries, duplicates = _deduplicate(parse_ocr_grid(items))
    warnings = [] if entries else ["Couldn't read any exams. Try a clearer timetable or enter exams manually."]
    if duplicates:
        warnings.append(f"Ignored {duplicates} repeated exam row(s) found across timetable pages.")
    return entries, warnings, len(items)
