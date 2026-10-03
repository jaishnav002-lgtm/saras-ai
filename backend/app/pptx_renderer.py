"""
PPTX visual slide parser.
Extracts structured layout, shapes, text formatting, images, and tables
so presentations can be rendered visually in the web client.
"""
import base64
import os
from typing import Any, Dict, List, Optional
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN


def _get_color_hex(color_format) -> Optional[str]:
    """Safely extracts RGB hex string from a ColorFormat object."""
    try:
        if color_format and hasattr(color_format, "rgb") and color_format.rgb:
            return f"#{color_format.rgb}"
    except Exception:
        pass
    return None


def _get_fill_color(shape) -> Optional[str]:
    """Safely extracts solid fill color from shape."""
    try:
        if hasattr(shape, "fill") and shape.fill and shape.fill.type == 1:  # 1 is MSO_FILL.SOLID
            return _get_color_hex(shape.fill.fore_color)
    except Exception:
        pass
    return None


def _get_line_color(shape) -> Optional[str]:
    """Safely extracts border/line color from shape."""
    try:
        if hasattr(shape, "line") and shape.line:
            return _get_color_hex(shape.line.color)
    except Exception:
        pass
    return None


def _get_alignment_str(alignment) -> str:
    if alignment == PP_ALIGN.CENTER:
        return "center"
    if alignment == PP_ALIGN.RIGHT:
        return "right"
    if alignment == PP_ALIGN.JUSTIFY:
        return "justify"
    return "left"


def _parse_paragraphs(text_frame) -> List[Dict[str, Any]]:
    paragraphs_data = []
    for para in text_frame.paragraphs:
        alignment = _get_alignment_str(para.alignment)
        level = getattr(para, "level", 0) or 0

        runs_data = []
        for run in para.runs:
            font_size = None
            try:
                if run.font and run.font.size:
                    font_size = round(run.font.size.pt, 1)
            except Exception:
                pass

            font_name = None
            try:
                if run.font and run.font.name:
                    font_name = run.font.name
            except Exception:
                pass

            color = None
            try:
                if run.font and run.font.color:
                    color = _get_color_hex(run.font.color)
            except Exception:
                pass

            runs_data.append({
                "text": run.text,
                "bold": bool(run.font.bold) if run.font else False,
                "italic": bool(run.font.italic) if run.font else False,
                "underline": bool(run.font.underline) if run.font else False,
                "font_size": font_size,
                "font_name": font_name,
                "color": color,
            })

        if not runs_data and para.text:
            runs_data.append({
                "text": para.text,
                "bold": False,
                "italic": False,
                "underline": False,
                "font_size": None,
                "font_name": None,
                "color": None,
            })

        paragraphs_data.append({
            "text": para.text,
            "level": level,
            "alignment": alignment,
            "runs": runs_data,
        })
    return paragraphs_data


def _parse_shape(shape, slide_width: int, slide_height: int) -> Optional[Dict[str, Any]]:
    try:
        left_pct = (shape.left / slide_width * 100) if slide_width else 0
        top_pct = (shape.top / slide_height * 100) if slide_height else 0
        width_pct = (shape.width / slide_width * 100) if slide_width else 0
        height_pct = (shape.height / slide_height * 100) if slide_height else 0
        rotation = getattr(shape, "rotation", 0) or 0

        base_elem = {
            "id": str(getattr(shape, "shape_id", id(shape))),
            "name": getattr(shape, "name", "Shape"),
            "left": round(left_pct, 2),
            "top": round(top_pct, 2),
            "width": round(width_pct, 2),
            "height": round(height_pct, 2),
            "rotation": rotation,
            "fill_color": _get_fill_color(shape),
            "border_color": _get_line_color(shape),
        }

        # 1. Pictures / Images
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
            try:
                img = shape.image
                ext = getattr(img, "ext", "png")
                b64 = base64.b64encode(img.blob).decode("utf-8")
                base_elem["type"] = "image"
                base_elem["src"] = f"data:image/{ext};base64,{b64}"
                base_elem["alt"] = getattr(shape, "name", "Slide Image")
                return base_elem
            except Exception:
                pass

        # 2. Tables
        if shape.has_table:
            base_elem["type"] = "table"
            table = shape.table
            rows_data = []
            for row in table.rows:
                row_cells = []
                for cell in row.cells:
                    cell_fill = _get_fill_color(cell)
                    cell_paras = _parse_paragraphs(cell.text_frame) if cell.text_frame else []
                    row_cells.append({
                        "text": cell.text,
                        "fill_color": cell_fill,
                        "paragraphs": cell_paras,
                    })
                rows_data.append(row_cells)
            base_elem["rows"] = rows_data
            return base_elem

        # 3. Text Boxes / Shapes with text
        if shape.has_text_frame:
            base_elem["type"] = "text"
            base_elem["paragraphs"] = _parse_paragraphs(shape.text_frame)
            return base_elem

        # 4. Auto shapes without text
        base_elem["type"] = "shape"
        return base_elem

    except Exception:
        return None


def parse_pptx_slides(path: str) -> Dict[str, Any]:
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")

    prs = Presentation(path)
    sw = prs.slide_width
    sh = prs.slide_height
    aspect_ratio = (sw / sh) if sh else 1.7777777777777777

    slides_data = []
    for idx, slide in enumerate(prs.slides, start=1):
        elements = []
        for shape in slide.shapes:
            # Handle group shapes
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP and hasattr(shape, "shapes"):
                for sub_shape in shape.shapes:
                    parsed_sub = _parse_shape(sub_shape, sw, sh)
                    if parsed_sub:
                        elements.append(parsed_sub)
            else:
                parsed = _parse_shape(shape, sw, sh)
                if parsed:
                    elements.append(parsed)

        # Slide background
        bg_color = None
        try:
            if hasattr(slide, "background") and slide.background:
                bg_color = _get_fill_color(slide.background)
        except Exception:
            pass

        slides_data.append({
            "slide_number": idx,
            "background_color": bg_color,
            "elements": elements,
        })

    return {
        "total_slides": len(slides_data),
        "aspect_ratio": round(aspect_ratio, 4),
        "slide_width": sw,
        "slide_height": sh,
        "slides": slides_data,
    }
