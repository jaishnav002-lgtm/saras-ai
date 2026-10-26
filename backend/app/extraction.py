"""
Text extraction for uploaded documents (PDF, PPTX, DOCX).
Preserves page/slide boundaries and individual page text units for RAG ingestion.
"""
from dataclasses import dataclass, field
import re
from typing import List, Optional


@dataclass
class DocumentPage:
    page_number: Optional[int]  # 1-based page count (pdf) or slide count (pptx); None for docx
    text: str


@dataclass
class ExtractionResult:
    full_text: str
    unit_count: int  # page count (pdf) or slide count (pptx); paragraphs for docx
    pages: List[DocumentPage] = field(default_factory=list)

    def has_extractable_text(self) -> bool:
        """
        Determines whether the document contains any genuine extractable text.
        Returns False if the document is empty or only contains scanned image placeholders.
        """
        if any(bool(p.text and p.text.strip()) for p in self.pages):
            return True
        # Fallback check stripping empty placeholders and headers
        cleaned = self.full_text
        for placeholder in ["[Empty page or image-only]", "[Empty slide or image-only]"]:
            cleaned = cleaned.replace(placeholder, "")
        cleaned = re.sub(r"--- (?:Page|Slide) \d+ ---", "", cleaned)
        return bool(cleaned.strip())


def extract_pdf(path: str) -> ExtractionResult:
    import fitz  # PyMuPDF

    doc = fitz.open(path)
    text_parts = []
    pages: List[DocumentPage] = []
    for idx, page in enumerate(doc, start=1):
        txt = page.get_text().strip()
        if txt:
            text_parts.append(f"--- Page {idx} ---\n{txt}")
            pages.append(DocumentPage(page_number=idx, text=txt))
        else:
            text_parts.append(f"--- Page {idx} ---\n[Empty page or image-only]")
            pages.append(DocumentPage(page_number=idx, text=""))
    page_count = len(doc)
    doc.close()
    return ExtractionResult(
        full_text="\n\n".join(text_parts),
        unit_count=page_count,
        pages=pages,
    )


def extract_pptx(path: str) -> ExtractionResult:
    from pptx import Presentation

    prs = Presentation(path)
    text_parts = []
    pages: List[DocumentPage] = []
    for idx, slide in enumerate(prs.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                slide_text.append(shape.text.strip())
        content = "\n".join(slide_text).strip()
        if content:
            text_parts.append(f"--- Slide {idx} ---\n{content}")
            pages.append(DocumentPage(page_number=idx, text=content))
        else:
            text_parts.append(f"--- Slide {idx} ---\n[Empty slide or image-only]")
            pages.append(DocumentPage(page_number=idx, text=""))
    return ExtractionResult(
        full_text="\n\n".join(text_parts),
        unit_count=len(prs.slides),
        pages=pages,
    )


def extract_docx(path: str) -> ExtractionResult:
    import docx

    d = docx.Document(path)
    paragraphs = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    full_text = "\n\n".join(paragraphs)
    pages: List[DocumentPage] = []
    if full_text:
        pages.append(DocumentPage(page_number=None, text=full_text))
    return ExtractionResult(
        full_text=full_text,
        unit_count=len(paragraphs),
        pages=pages,
    )


def extract_text(path: str, file_type: str) -> ExtractionResult:
    file_type = file_type.lower()
    if file_type == "pdf":
        return extract_pdf(path)
    if file_type == "pptx":
        return extract_pptx(path)
    if file_type == "docx":
        return extract_docx(path)
    raise ValueError(f"Unsupported file type: {file_type}")
