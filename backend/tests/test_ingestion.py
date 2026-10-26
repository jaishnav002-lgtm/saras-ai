import io
import os
import sys
import uuid
from pathlib import Path
from unittest.mock import MagicMock, patch

import fitz  # PyMuPDF
import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches, Pt
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.database import Base
from app.extraction import DocumentPage, ExtractionResult, extract_text
from app.models import Document, DocumentChunk, Subject, User
from app.services.exceptions import IngestionError, UnextractableDocumentError
from app.services.ingestion import (
    DocumentIngestionService,
    generate_stable_chunk_id,
    split_text_into_chunks,
    validate_chunk_settings,
)

# Test SQLite Database Setup
TEST_DB_URL = "sqlite:///./test_ingestion.db"
test_engine = create_engine(TEST_DB_URL, connect_args={"check_same_thread": False})
TestSession = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


@pytest.fixture(autouse=True)
def restore_dependency_overrides():
    """Keep route-level test overrides from leaking into other test modules."""
    from app.main import app

    original = dict(app.dependency_overrides)
    yield
    app.dependency_overrides.clear()
    app.dependency_overrides.update(original)


@pytest.fixture(autouse=True)
def mock_route_indexing():
    """Route tests exercise extraction/ingestion without a live NVIDIA call."""
    with patch("app.routers.documents.DocumentEmbeddingIndexer.index_chunks", return_value=0):
        yield


@pytest.fixture(scope="module", autouse=True)
def setup_test_database():
    Base.metadata.create_all(bind=test_engine)
    yield
    Base.metadata.drop_all(bind=test_engine)
    if os.path.exists("test_ingestion.db"):
        os.remove("test_ingestion.db")


@pytest.fixture
def db_session():
    session = TestSession()
    try:
        yield session
    finally:
        session.rollback()
        session.close()


@pytest.fixture
def sample_user_and_subject(db_session):
    unique_suffix = uuid.uuid4().hex[:8]
    user = User(
        name="Test Student",
        email=f"student_{unique_suffix}@example.com",
        password_hash="fakehash",
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)

    subject = Subject(
        user_id=user.id,
        name="Operating Systems",
        description="OS Concepts",
    )
    db_session.add(subject)
    db_session.commit()
    db_session.refresh(subject)

    return user, subject


# ---------------------------------------------------------------------------
# 1. Text Extraction & Page-Number Preservation
# ---------------------------------------------------------------------------

def test_pdf_extraction_page_number_preservation(tmp_path):
    """Verify PyMuPDF extraction preserves 1-based page numbers and boundaries."""
    pdf_path = tmp_path / "multipage.pdf"
    doc = fitz.open()

    page1 = doc.new_page()
    page1.insert_text((50, 72), "First page content about process scheduling.")

    page2 = doc.new_page()
    page2.insert_text((50, 72), "Second page content about virtual memory and paging.")

    doc.save(str(pdf_path))
    doc.close()

    result = extract_text(str(pdf_path), "pdf")

    assert result.unit_count == 2
    assert len(result.pages) == 2
    assert result.pages[0].page_number == 1
    assert "process scheduling" in result.pages[0].text
    assert result.pages[1].page_number == 2
    assert "virtual memory" in result.pages[1].text
    assert result.has_extractable_text() is True


def test_pptx_extraction_slide_number_preservation(tmp_path):
    """Verify PPTX extraction preserves slide numbers."""
    pptx_path = tmp_path / "presentation.pptx"
    prs = Presentation()

    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    tb1 = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    tb1.text_frame.text = "Slide 1: Threads and Concurrency."

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    tb2 = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    tb2.text_frame.text = "Slide 2: Deadlocks and Mutexes."

    prs.save(str(pptx_path))

    result = extract_text(str(pptx_path), "pptx")

    assert result.unit_count == 2
    assert len(result.pages) == 2
    assert result.pages[0].page_number == 1
    assert "Threads and Concurrency" in result.pages[0].text
    assert result.pages[1].page_number == 2
    assert "Deadlocks and Mutexes" in result.pages[1].text
    assert result.has_extractable_text() is True


def test_docx_extraction_without_physical_pages(tmp_path):
    """Verify DOCX extraction assigns page_number=None for flow documents."""
    docx_path = tmp_path / "document.docx"
    doc = DocxDocument()
    doc.add_paragraph("Paragraph 1 on CPU scheduling.")
    doc.add_paragraph("Paragraph 2 on synchronization primitives.")
    doc.save(str(docx_path))

    result = extract_text(str(docx_path), "docx")

    assert result.unit_count == 2
    assert len(result.pages) == 1
    assert result.pages[0].page_number is None
    assert "CPU scheduling" in result.pages[0].text
    assert "synchronization primitives" in result.pages[0].text
    assert result.has_extractable_text() is True


# ---------------------------------------------------------------------------
# 2. Chunk Size, Overlap, and Chunk Indices
# ---------------------------------------------------------------------------

def test_chunk_settings_validation():
    """Verify validation of chunk_size and chunk_overlap."""
    # Valid
    validate_chunk_settings(chunk_size=500, chunk_overlap=100)
    validate_chunk_settings(chunk_size=100, chunk_overlap=0)

    # Invalid chunk_size <= 0
    with pytest.raises(ValueError, match="positive integer"):
        validate_chunk_settings(chunk_size=0, chunk_overlap=0)
    with pytest.raises(ValueError, match="positive integer"):
        validate_chunk_settings(chunk_size=-10, chunk_overlap=0)

    # Invalid negative overlap
    with pytest.raises(ValueError, match="cannot be negative"):
        validate_chunk_settings(chunk_size=100, chunk_overlap=-1)

    # Invalid overlap >= chunk_size
    with pytest.raises(ValueError, match="strictly less"):
        validate_chunk_settings(chunk_size=100, chunk_overlap=100)
    with pytest.raises(ValueError, match="strictly less"):
        validate_chunk_settings(chunk_size=100, chunk_overlap=150)


def test_split_text_into_chunks_size_and_overlap():
    """Verify text is chunked according to max size and overlap."""
    text = (
        "Operating systems manage hardware resources. "
        "A process is a program in execution. "
        "Threads share the same address space. "
        "Virtual memory allows large address spaces. "
        "File systems provide persistent storage."
    )
    chunks = split_text_into_chunks(text, chunk_size=80, chunk_overlap=20)

    assert len(chunks) > 1
    for chunk in chunks:
        assert len(chunk) <= 80
        assert len(chunk) > 0


def test_chunk_pages_sequential_indices():
    """Verify chunk indices are monotonically incrementing across all pages."""
    service = DocumentIngestionService()
    pages = [
        DocumentPage(page_number=1, text="Page 1 sentence A. Page 1 sentence B."),
        DocumentPage(page_number=2, text="Page 2 sentence C. Page 2 sentence D."),
    ]

    chunk_dicts = service.chunk_pages(pages, chunk_size=30, chunk_overlap=5)

    assert len(chunk_dicts) >= 2
    for i, c in enumerate(chunk_dicts):
        assert c["chunk_index"] == i
        assert len(c["content"]) > 0
        assert c["page_number"] in (1, 2)
        assert len(c["content_hash"]) == 64


# ---------------------------------------------------------------------------
# 3. Preservation of Document, Subject, Owner, Filename, Page, & Status
# ---------------------------------------------------------------------------

def test_ingestion_preserves_full_metadata(db_session, sample_user_and_subject, tmp_path):
    """Verify chunk record retains document_id, subject_id, owner_id, filename, page, status, hash."""
    user, subject = sample_user_and_subject

    # Create dummy file
    file_path = tmp_path / "os_notes.txt"
    file_path.write_text("Page content about memory management.")

    document = Document(
        subject_id=subject.id,
        filename="os_notes.pdf",
        file_type="pdf",
        storage_path=str(file_path),
        status="uploaded",
    )
    db_session.add(document)
    db_session.commit()
    db_session.refresh(document)

    extraction = ExtractionResult(
        full_text="Memory management page 1 content.\n\nVirtual memory page 2 content.",
        unit_count=2,
        pages=[
            DocumentPage(page_number=1, text="Memory management page 1 content."),
            DocumentPage(page_number=2, text="Virtual memory page 2 content."),
        ],
    )

    service = DocumentIngestionService()
    chunks = service.ingest_document(
        document=document,
        db=db_session,
        extraction_result=extraction,
        owner_id=user.id,
        chunk_size=100,
        chunk_overlap=20,
    )

    assert len(chunks) == 2
    assert document.status == "processed"
    assert document.error_message is None

    # Inspect first chunk
    c1 = chunks[0]
    assert c1.document_id == document.id
    assert c1.subject_id == subject.id
    assert c1.owner_id == user.id
    assert c1.filename == "os_notes.pdf"
    assert c1.page_number == 1
    assert c1.chunk_index == 0
    assert c1.status == "ingested"
    assert "Memory management" in c1.content
    assert len(c1.content_hash) == 64
    assert c1.token_count > 0

    # Inspect second chunk
    c2 = chunks[1]
    assert c2.page_number == 2
    assert c2.chunk_index == 1
    assert "Virtual memory" in c2.content


# ---------------------------------------------------------------------------
# 4. Repeated Ingestion & Retry Idempotency
# ---------------------------------------------------------------------------

def test_repeated_ingestion_idempotency(db_session, sample_user_and_subject, tmp_path):
    """Reprocessing the same document must not duplicate chunks and must preserve stable IDs."""
    user, subject = sample_user_and_subject
    file_path = tmp_path / "idempotent_doc.txt"
    file_path.write_text("Persistent content for idempotency test.")

    document = Document(
        subject_id=subject.id,
        filename="idempotent.docx",
        file_type="docx",
        storage_path=str(file_path),
        status="uploaded",
    )
    db_session.add(document)
    db_session.commit()
    db_session.refresh(document)

    extraction = ExtractionResult(
        full_text="Sentence 1 for testing. Sentence 2 for testing.",
        unit_count=1,
        pages=[DocumentPage(page_number=None, text="Sentence 1 for testing. Sentence 2 for testing.")],
    )

    service = DocumentIngestionService()

    # First ingestion run
    first_run = service.ingest_document(
        document=document,
        db=db_session,
        extraction_result=extraction,
        owner_id=user.id,
        chunk_size=50,
        chunk_overlap=10,
    )
    first_chunk_ids = [c.id for c in first_run]
    first_count = len(first_run)
    assert first_count >= 1

    # Verify count in database
    db_chunks_1 = (
        db_session.query(DocumentChunk)
        .filter(DocumentChunk.document_id == document.id)
        .all()
    )
    assert len(db_chunks_1) == first_count

    # Second ingestion run (retry/reprocess)
    second_run = service.ingest_document(
        document=document,
        db=db_session,
        extraction_result=extraction,
        owner_id=user.id,
        chunk_size=50,
        chunk_overlap=10,
    )
    second_chunk_ids = [c.id for c in second_run]

    # Verify no duplicate chunks in database
    db_chunks_2 = (
        db_session.query(DocumentChunk)
        .filter(DocumentChunk.document_id == document.id)
        .all()
    )
    assert len(db_chunks_2) == first_count
    # Verify deterministic, stable chunk identity
    assert first_chunk_ids == second_chunk_ids


def test_stable_chunk_id_determinism():
    """Verify generate_stable_chunk_id returns identical UUIDs for the same document and index."""
    doc_id = "doc-12345-abcde"
    id1 = generate_stable_chunk_id(doc_id, 0)
    id2 = generate_stable_chunk_id(doc_id, 0)
    id3 = generate_stable_chunk_id(doc_id, 1)

    assert id1 == id2
    assert id1 != id3


# ---------------------------------------------------------------------------
# 5. Empty / Unextractable Documents & Failure Handling
# ---------------------------------------------------------------------------

def test_unextractable_document_raises_and_updates_status(db_session, sample_user_and_subject, tmp_path):
    """An empty or image-only document must raise UnextractableDocumentError and mark status='failed'."""
    user, subject = sample_user_and_subject
    file_path = tmp_path / "empty.pdf"
    file_path.write_text("")

    document = Document(
        subject_id=subject.id,
        filename="empty.pdf",
        file_type="pdf",
        storage_path=str(file_path),
        status="uploaded",
    )
    db_session.add(document)
    db_session.commit()
    db_session.refresh(document)

    # Simulate extraction with no text (e.g. scanned image with empty placeholder)
    empty_extraction = ExtractionResult(
        full_text="--- Page 1 ---\n[Empty page or image-only]",
        unit_count=1,
        pages=[DocumentPage(page_number=1, text="")],
    )
    assert empty_extraction.has_extractable_text() is False

    service = DocumentIngestionService()
    with pytest.raises(UnextractableDocumentError) as exc_info:
        service.ingest_document(
            document=document,
            db=db_session,
            extraction_result=empty_extraction,
            owner_id=user.id,
        )

    assert "no extractable text" in str(exc_info.value).lower()
    db_session.refresh(document)
    assert document.status == "failed"
    assert "no extractable text" in document.error_message.lower()

    # Confirm 0 chunks were stored
    stored_chunks = (
        db_session.query(DocumentChunk)
        .filter(DocumentChunk.document_id == document.id)
        .all()
    )
    assert len(stored_chunks) == 0


def test_whitespace_only_document(db_session, sample_user_and_subject, tmp_path):
    """Whitespace-only document is rejected as unextractable."""
    user, subject = sample_user_and_subject
    file_path = tmp_path / "spaces.docx"
    file_path.write_text("   \n\t  ")

    document = Document(
        subject_id=subject.id,
        filename="spaces.docx",
        file_type="docx",
        storage_path=str(file_path),
        status="uploaded",
    )
    db_session.add(document)
    db_session.commit()
    db_session.refresh(document)

    whitespace_extraction = ExtractionResult(
        full_text="    \n   ",
        unit_count=1,
        pages=[DocumentPage(page_number=None, text="   \n\t  ")],
    )

    service = DocumentIngestionService()
    with pytest.raises(UnextractableDocumentError):
        service.ingest_document(
            document=document,
            db=db_session,
            extraction_result=whitespace_extraction,
            owner_id=user.id,
        )

    db_session.refresh(document)
    assert document.status == "failed"


def test_ingestion_transaction_rollback_on_db_error(db_session, sample_user_and_subject, tmp_path):
    """Database insertion errors must roll back the transaction and mark status='failed'."""
    user, subject = sample_user_and_subject
    file_path = tmp_path / "valid.txt"
    file_path.write_text("Valid text.")

    document = Document(
        subject_id=subject.id,
        filename="valid.txt",
        file_type="pdf",
        storage_path=str(file_path),
        status="uploaded",
    )
    db_session.add(document)
    db_session.commit()
    db_session.refresh(document)

    extraction = ExtractionResult(
        full_text="Valid text content.",
        unit_count=1,
        pages=[DocumentPage(page_number=1, text="Valid text content.")],
    )

    service = DocumentIngestionService()

    # Mock db.add to simulate a database constraint/disk error
    with patch.object(db_session, "add", side_effect=Exception("Disk full or constraint error")):
        with pytest.raises(IngestionError, match="Ingestion storage failed"):
            service.ingest_document(
                document=document,
                db=db_session,
                extraction_result=extraction,
                owner_id=user.id,
            )

    db_session.refresh(document)
    assert document.status == "failed"
    assert "storage failed" in document.error_message.lower()


def test_helper_get_pending_chunks(db_session, sample_user_and_subject, tmp_path):
    """Verify get_pending_chunks and get_document_chunks helpers for subsequent RAG embedding stage."""
    user, subject = sample_user_and_subject
    file_path = tmp_path / "pending_test.txt"
    file_path.write_text("Content")

    document = Document(
        subject_id=subject.id,
        filename="pending.docx",
        file_type="docx",
        storage_path=str(file_path),
        status="uploaded",
    )
    db_session.add(document)
    db_session.commit()
    db_session.refresh(document)

    extraction = ExtractionResult(
        full_text="Testing chunk lookups.",
        unit_count=1,
        pages=[DocumentPage(page_number=None, text="Testing chunk lookups.")],
    )

    service = DocumentIngestionService()
    chunks = service.ingest_document(
        document=document,
        db=db_session,
        extraction_result=extraction,
        owner_id=user.id,
    )

    all_chunks = service.get_document_chunks(document.id, db_session)
    assert len(all_chunks) == len(chunks)

    pending_chunks = service.get_pending_chunks(document.id, db_session)
    assert len(pending_chunks) == len(chunks)
    assert pending_chunks[0].status == "ingested"


def test_upload_route_e2e_creates_chunks_and_preserves_metadata():
    """Verify that the FastAPI upload endpoint populates document_chunks with owner and metadata."""
    from fastapi.testclient import TestClient
    from app.main import app
    from app.database import get_db

    def override_get_db_for_route():
        session = TestSession()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_get_db_for_route
    client = TestClient(app)

    # 1. Sign up user
    email = f"ingest_user_{uuid.uuid4().hex[:8]}@example.com"
    signup_res = client.post("/api/auth/signup", json={"name": "Ingest User", "email": email, "password": "password123"})
    assert signup_res.status_code == 200
    token = signup_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 2. Create Subject
    subj_res = client.post("/api/subjects", json={"name": "Compiler Design"}, headers=headers)
    assert subj_res.status_code == 200
    subject_id = subj_res.json()["id"]

    # 3. Create DOCX with content
    doc = DocxDocument()
    doc.add_paragraph("Lexical analysis converts stream of characters into tokens.")
    doc.add_paragraph("Syntax analysis builds an abstract syntax tree from tokens.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    # 4. Upload document
    upload_res = client.post(
        f"/api/subjects/{subject_id}/documents",
        files={"file": ("lexer.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=headers,
    )
    assert upload_res.status_code == 200
    doc_body = upload_res.json()
    assert doc_body["status"] == "processed"
    doc_id = doc_body["id"]

    # 5. Query DB directly to verify DocumentChunk rows
    db = TestSession()
    try:
        chunks = db.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).all()
        assert len(chunks) >= 1
        for chunk in chunks:
            assert chunk.document_id == doc_id
            assert chunk.subject_id == subject_id
            assert chunk.filename == "lexer.docx"
            assert chunk.status == "ingested"
            assert len(chunk.content_hash) == 64
            assert "Lexical analysis" in chunk.content or "Syntax analysis" in chunk.content
    finally:
        db.close()


def test_upload_route_empty_docx_reports_clear_actionable_error():
    """Verify that uploading an empty docx reports status='failed' with actionable error."""
    from fastapi.testclient import TestClient
    from app.main import app
    from app.database import get_db

    def override_get_db_for_route():
        session = TestSession()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_get_db_for_route
    client = TestClient(app)

    # 1. Sign up user
    email = f"empty_user_{uuid.uuid4().hex[:8]}@example.com"
    signup_res = client.post("/api/auth/signup", json={"name": "Empty User", "email": email, "password": "password123"})
    token = signup_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 2. Create Subject
    subj_res = client.post("/api/subjects", json={"name": "Empty Course"}, headers=headers)
    subject_id = subj_res.json()["id"]

    # 3. Create completely empty DOCX (no paragraphs)
    doc = DocxDocument()
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    # 4. Upload document
    upload_res = client.post(
        f"/api/subjects/{subject_id}/documents",
        files={"file": ("blank.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=headers,
    )
    assert upload_res.status_code == 200
    doc_body = upload_res.json()
    assert doc_body["status"] == "failed"
    assert "no extractable text" in doc_body["error_message"].lower()

    # 5. Verify 0 chunks created
    doc_id = doc_body["id"]
    db = TestSession()
    try:
        chunks = db.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).all()
        assert len(chunks) == 0
    finally:
        db.close()
